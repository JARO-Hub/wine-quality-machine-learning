from __future__ import annotations

import argparse
import csv
import json
import os
import re
from collections.abc import Callable
from datetime import date
from pathlib import Path
from typing import cast

import matplotlib

matplotlib.use("Agg")
import matplotlib.dates as mdates
import matplotlib.pyplot as plt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from matplotlib.ticker import Formatter

from wine_quality.shared.domain.types import JsonObject, JsonValue


def _object(value: JsonValue) -> JsonObject:
    if not isinstance(value, dict):
        raise ValueError("Se esperaba un objeto del informe.")
    return value


def _objects(value: JsonValue) -> list[JsonObject]:
    if not isinstance(value, list):
        raise ValueError("Se esperaba una lista del informe.")
    return [_object(item) for item in value]


def _number(value: JsonValue, places: int = 4) -> str:
    if value is None:
        return "indefinido"
    if not isinstance(value, (int, float)):
        raise ValueError("Se esperaba una métrica numérica.")
    return f"{value:.{places}f}"


def _table(headers: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    lines.extend("| " + " | ".join(row) + " |" for row in rows)
    return "\n".join(lines)


def _figures(report: JsonObject, directory: Path) -> tuple[Path, Path]:
    with (directory / "series.csv").open(encoding="utf-8", newline="") as source:
        rows = list(csv.DictReader(source))
    dates = [date.fromisoformat(row["date"]) for row in rows]
    cumulative = [float(row["cumulative"]) for row in rows]
    changes = [float(row["daily_change"]) if row["daily_change"] else float("nan") for row in rows]
    split = _object(report["split"])
    test_start = date.fromisoformat(str(_object(split["test"])["first_date"]))
    plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
    figure, axes = plt.subplots(2, 1, figsize=(9, 5.5), sharex=True, constrained_layout=True)
    axes[0].plot(dates, cumulative, color="#23688e")
    axes[0].set_ylabel("Casos acumulados")
    axes[1].plot(dates, changes, color="#23688e", linewidth=0.9)
    axes[1].set_ylabel("Variación registrada/día")
    axes[1].set_xlabel("Fecha del registro")
    for axis in axes:
        axis.axvspan(test_start, dates[-1], color="#b76321", alpha=0.12, label="Periodo de prueba")
        axis.grid(alpha=0.2)
    axes[0].legend(loc="upper left")
    date_formatter = cast(Callable[[str], Formatter], mdates.DateFormatter)
    axes[1].xaxis.set_major_formatter(date_formatter("%Y-%m"))
    history_path = directory / "series.png"
    figure.savefig(history_path, dpi=180)
    plt.close(figure)

    models = sorted(
        _objects(report["models"]), key=lambda m: cast(float, _object(m["cv"])["rmse_mean"])
    )
    figure, axes = plt.subplots(2, 1, figsize=(9, 6), constrained_layout=True)
    axes[0].barh(
        [str(model["label"]) for model in models],
        [cast(float, _object(model["cv"])["rmse_mean"]) for model in models],
        xerr=[cast(float, _object(model["cv"])["rmse_std"]) for model in models],
        color="#23688e",
        capsize=3,
    )
    axes[0].invert_yaxis()
    axes[0].set_xlabel("RMSE CV ± desviación entre folds (no intervalo de confianza)")
    points = _objects(_object(_object(report["test"])["winner"])["predictions"])
    test_dates = [date.fromisoformat(str(point["date"])) for point in points]
    axes[1].plot(
        test_dates, [point["actual"] for point in points], label="Observado", color="#23688e"
    )
    axes[1].plot(
        test_dates,
        [point["predicted"] for point in points],
        label="Modelo elegido",
        color="#b76321",
    )
    axes[1].set_ylabel("Variación registrada/día")
    axes[1].set_xlabel("Fecha objetivo en prueba")
    axes[1].xaxis.set_major_formatter(date_formatter("%Y-%m-%d"))
    axes[1].legend()
    for axis in axes:
        axis.grid(alpha=0.2)
    evaluation_path = directory / "evaluation.png"
    figure.savefig(evaluation_path, dpi=180)
    plt.close(figure)
    return history_path, evaluation_path


def _markdown(report: JsonObject, images: tuple[Path, Path], report_directory: Path) -> str:
    data = _object(report["data"])
    experiment = _object(report["experiment"])
    split = _object(report["split"])
    train, test_split = _object(split["train"]), _object(split["test"])
    models = sorted(
        _objects(report["models"]), key=lambda m: cast(float, _object(m["cv"])["rmse_mean"])
    )
    winner = next(model for model in models if model["id"] == report["winner_id"])
    test = _object(report["test"])
    winner_test = _object(test["winner"])
    metrics = _object(winner_test["metrics"])
    country = str(data["country"])
    persistence_cv = next(
        baseline for baseline in _objects(report["baseline_cv"]) if baseline["lag"] == 1
    )
    cv_relation = (
        "menor"
        if cast(float, persistence_cv["rmse_mean"])
        < cast(float, _object(winner["cv"])["rmse_mean"])
        else "mayor o igual"
    )
    parts: list[str] = [
        f"# Informe integrado de COVID: {country}",
        "Grupo 15 · Machine Learning · Caso 04 · 30 de septiembre de 2026",
        "## 1. Pregunta, integración y resultado principal",
        f"Este caso estudia si el historial de casos confirmados registrados en {country} "
        "permite estimar la variación del registro del día siguiente. Integra preparación, "
        "comparación de regresores, una red neuronal y evaluación temporal. Comparte "
        "infraestructura con el proyecto de vinos, pero tiene datos, respuesta, particiones "
        "e informe propios. No se mezclan unidades epidemiológicas con mediciones de vino.",
        f"El modelo elegido exclusivamente por validación fue {winner['label']}, con RMSE "
        f"medio {_number(_object(winner['cv'])['rmse_mean'])}. En la prueba final obtuvo "
        f"RMSE {_number(metrics['rmse'])} y MAE {_number(metrics['mae'])}. Estas cifras "
        "miden errores en la variación diaria registrada, no aciertos diagnósticos. "
        "La conclusión frente a referencias se desarrolla en la sección 7.",
        "## 2. Fuente, estructura e integridad de los datos",
        "El archivo aportado se llama time_series_covid_19_confirmed.csv. Su estructura "
        "coincide con las series globales JHU CSSE: provincia/estado, país/región, latitud, "
        "longitud y columnas por fecha. La documentación de series [1], apartado Time "
        "series summary, describe esa organización. No se comprobó igualdad byte por byte "
        "con una revisión remota ni se sustituyó el archivo del usuario.",
        "El README JHU [2], apartado Terms of Use, atribuye el dataset a JHU CSSE y declara "
        "CC BY 4.0. La referencia bibliográfica solicitada allí es Dong, Du y Gardner "
        "(2020), DOI 10.1016/S1473-3099(20)30120-1. Las fuentes web no tienen paginación: "
        "se localizan por apartado. La licencia del dataset no se extiende al código "
        "o al informe del grupo.",
        _table(
            ["Descripción", "Valor"],
            [
                ["Filas geográficas originales", str(data["rows_raw"])],
                ["Países/regiones etiquetados en el CSV", str(data["countries_in_csv"])],
                ["Fechas diarias", str(data["date_count"])],
                ["Periodo", f"{data['first_date']} a {data['last_date']}"],
                ["Duplicados geográficos exactos eliminados", str(data["duplicates_removed"])],
                ["Filas originales utilizadas para el país", str(data["source_row_ids"])],
                ["Acumulado final del país", _number(data["last_cumulative"], 0)],
                ["Revisiones diarias negativas en el país", str(data["negative_revision_count"])],
                ["Muestras supervisadas", str(data["supervised_count"])],
            ],
        ),
        f"Huella SHA256 de los bytes del CSV: {data['input_sha256']}. Esta huella "
        "identifica la copia efectiva, incluidos sus finales de línea. Cada fecha "
        "objetivo conserva además su posición original entre las columnas de fechas.",
        "## 3. Preparación: del acumulado a una respuesta diaria",
        "Se comprueban fechas consecutivas, valores acumulados enteros no negativos y "
        "ausencia de faltantes en los conteos. Un nombre de provincia vacío puede "
        "representar una fila nacional; no se imputa como una medición faltante. "
        "Se quitan únicamente filas geográficas completamente iguales, conservando "
        "la primera y su identificador. Filas de la misma provincia y país con "
        "valores distintos producen un error antes de agregar.",
        "Sea C_t el acumulado del país en la fecha t. Primero se suman las filas "
        "correspondientes al país. Después se calcula d_t = C_t − C_(t−1). Esta "
        "diferencia es la respuesta diaria registrada. El primer día no tiene un "
        "antecedente y su diferencia queda indefinida; no se inventa un cero. "
        "Una diferencia negativa se conserva como revisión. La sección Data "
        "modification records de JHU [3] documenta cambios retrospectivos; esta "
        "lectura motiva distinguir registros e infecciones, pero la política de "
        "conservar las diferencias es nuestra decisión.",
        "Para pronosticar d_t al cierre de t−1 se construyen 18 entradas: los "
        "14 valores d_(t−1), …, d_(t−14); las medias de los últimos 7 y 14 días; "
        "y seno y coseno del día de semana objetivo. La media de 7 días suma "
        "d_(t−1) hasta d_(t−7) y divide entre 7. La ventana no está centrada y "
        "no contiene d_t. Para el calendario, k es el día de semana con lunes=0: "
        "se usan sen(2πk/7) y cos(2πk/7). La fecha futura es conocida al emitir "
        "el pronóstico y no aporta casos futuros.",
        "Se excluyen 15 fechas iniciales por la diferencia indefinida y la necesidad "
        "de 14 antecedentes. Se conservan los ceros iniciales, los extremos y "
        "los días distintos con valores iguales. No se recorta ni redondea la "
        "respuesta. Estudiar la diferencia evita que una tendencia acumulativa "
        "suave domine la evaluación; no elimina por sí solo los cambios de régimen.",
    ]
    caption = "Figura 1. Acumulados y variaciones del registro; sombreado: periodo de prueba."
    relative = Path(os.path.relpath(images[0], report_directory))
    parts.extend([f"![{caption}]({relative.as_posix()})", caption])
    parts.extend(
        [
            "## 4. Protocolo temporal y control de información futura",
            f"Se reservaron las últimas {test_split['count']} muestras para prueba: "
            f"{test_split['first_date']} a {test_split['last_date']}. Entrenamiento "
            f"contiene {train['count']} muestras, de {train['first_date']} a "
            f"{train['last_date']}. La cantidad de prueba se obtiene redondeando hacia "
            "arriba el 20 % de las muestras supervisadas. Esta fracción se fijó antes "
            "de calcular resultados y no depende de cuál modelo salga favorecido.",
            "TimeSeriesSplit [4], apartado Parameters, implementa cortes ordenados "
            "y permite que el entrenamiento crezca. Elegimos cinco folds expansivos "
            "sin barajar. Toda fecha objetivo de ajuste precede a las de validación. "
            "Usamos gap=0 porque el horizonte es un día y d_(t−1) se conoce al "
            "pronosticar d_t; compartir historia entre ventanas no constituye "
            "por sí mismo fuga. Esa decisión no sería automáticamente válida "
            "para respuestas que agregaran varios días futuros.",
            _table(
                ["Fold", "Ajuste: fechas (n)", "Validación: fechas (n)"],
                [
                    [
                        str(i + 1),
                        f"{_object(fold['train'])['first_date']} a "
                        f"{_object(fold['train'])['last_date']} "
                        f"({_object(fold['train'])['count']})",
                        f"{_object(fold['validation'])['first_date']} a "
                        f"{_object(fold['validation'])['last_date']} "
                        f"({_object(fold['validation'])['count']})",
                    ]
                    for i, fold in enumerate(_objects(split["cv_folds"]))
                ],
            ),
            "Cada fold aprende sus propias medias y escalas. StandardScaler se "
            "encuentra dentro del Pipeline para entradas de modelos lineales, SVR "
            "y RNA. Los árboles reciben entradas sin escalar. Todos los candidatos "
            "transforman y con StandardScaler mediante TransformedTargetRegressor "
            "[5], apartado Parameters: el modelo aprende en la escala transformada "
            "y la transformación inversa devuelve casos por día. Esa escala también "
            "se aprende únicamente con las respuestas de ajuste de cada fold.",
            "Se selecciona la menor media de RMSE por fold; los empates exactos "
            "siguen el orden del catálogo. El elegido se reajusta con todo "
            "entrenamiento, manteniendo fijos sus parámetros en prueba. Para cada "
            "fecha de prueba se incorporan los registros reales ya observados "
            "hasta el día anterior. Esto evalúa predicciones sucesivas a un día, "
            "no una trayectoria de todo el periodo emitida desde un único origen. "
            "Solo el elegido y dos referencias fijadas se evalúan en prueba.",
            "## 5. Modelos y relación con sus operaciones",
            "OLS relaciona las entradas con una combinación lineal: b + Σ w_j x_j. "
            "b es el intercepto, w_j el coeficiente de la entrada j y la suma "
            "recorre las 18 entradas. Aprende minimizando la suma de errores "
            "cuadrados. Ridge añade alpha × Σ w_j² para penalizar coeficientes "
            "grandes. Se comparan alpha 0.1, 1, 10 y 100. Un coeficiente no "
            "representa un efecto causal de una intervención sanitaria.",
            "CART divide entradas para reducir error cuadrático y predice la "
            "media de las respuestas de su hoja. Se prueban profundidades 3, 5 "
            "y sin límite, y mínimos de 5 o 15 muestras por hoja. El bosque "
            "promedia 160 árboles ajustados con remuestreo con reemplazo. Compara "
            "profundidad 8 o sin límite, hojas de 2 o 5 muestras y candidatos "
            "de variables 1.0 o sqrt. Estas grillas se heredaron del catálogo "
            "académico; no se presentan como óptimas para epidemiología.",
            "SVR utiliza una tolerancia epsilon: los residuos dentro de ese "
            "margen no aportan pérdida epsilon-insensible. C controla la "
            "penalización de residuos que lo exceden. El kernel lineal representa "
            "una relación lineal; RBF permite relaciones no lineales mediante "
            "K(u,v) = exp(−gamma × ||u−v||²). u y v son vectores de entradas "
            "estandarizadas y gamma controla la caída de similitud con distancia. "
            "C y epsilon están en la escala de la respuesta estandarizada del "
            "fold, no en casos originales. Las grillas completas están en report.json.",
            "La RNA fija tiene 18 entradas, 8 neuronas ocultas tanh y una salida "
            "lineal. Para una fila estandarizada u, calcula h = tanh(uW₁+b₁) y "
            "z = hW₂+b₂. W₁ tiene forma 18×8, b₁ contiene 8 sesgos, W₂ tiene "
            "forma 8×1 y b₂ un sesgo. Son 161 parámetros: 144+8+8+1. Si mu_y "
            "y s_y son la media y escala de la respuesta de ajuste, la predicción "
            "final es mu_y+s_y×z. La salida es continua y puede ser negativa.",
            "La red utiliza L-BFGS, alpha=0.001, max_iter=2000, max_fun=50000 y "
            "semilla 42. Backpropagation calcula gradientes y L-BFGS actualiza "
            "parámetros con un método cuasi Newton. No es el SGD ni la red 11→8→1 "
            "del caso de vinos. Se fijó un único candidato neuronal antes de "
            "entrenar. Los 43 ajustes de hiperparámetros del catálogo previo "
            "más esta configuración suman 44 configuraciones en cinco folds.",
            "## 6. Métricas y resultados de validación",
            "Para n fechas evaluadas, y_i es la variación observada y ŷ_i la "
            "predicción. El residuo es r_i = y_i−ŷ_i. MAE promedia |r_i|; "
            "MSE promedia r_i²; RMSE es la raíz de MSE. MAE y RMSE se expresan "
            "en casos registrados por día, MSE en su cuadrado. R² compara la "
            "suma de errores cuadrados con la variación respecto a la media "
            "observada del conjunto evaluado: 1−Σr_i²/Σ(y_i−ȳ)². Puede ser "
            "negativo y queda indefinido si todas las respuestas son iguales.",
            _table(
                ["Modelo", "RMSE CV", "DE entre folds", "MAE CV", "Avisos"],
                [
                    [
                        str(model["label"]),
                        _number(_object(model["cv"])["rmse_mean"]),
                        _number(_object(model["cv"])["rmse_std"]),
                        _number(_object(model["cv"])["mae_mean"]),
                        str(len(_objects(model["fit_warnings"]))),
                    ]
                    for model in models
                ],
            ),
            "DE describe la dispersión de los cinco RMSE; no es un intervalo "
            "de confianza. Los folds corresponden a épocas distintas y sus "
            "errores no prueban superioridad estadística. Elegir entre "
            "configuraciones también puede favorecer la cifra de validación "
            "seleccionada; por eso la prueba se mantiene fuera de la elección.",
            _table(
                ["Referencia fija", "RMSE CV", "DE entre folds"],
                [
                    [
                        "Persistencia de un día"
                        if baseline["lag"] == 1
                        else "Persistencia semanal",
                        _number(baseline["rmse_mean"]),
                        _number(baseline["rmse_std"]),
                    ]
                    for baseline in _objects(report["baseline_cv"])
                ],
            ),
            "Las referencias predicen d_(t−1) y d_(t−7), respectivamente. No "
            "necesitan entrenar ni escalar. Permiten comprobar si la complejidad "
            "adicional mejora reglas que aprovechan continuidad o calendario "
            "semanal. No participaron en la selección del candidato aprendido.",
            f"La persistencia de un día obtuvo RMSE CV "
            f"{_number(persistence_cv['rmse_mean'])}, frente a "
            f"{_number(_object(winner['cv'])['rmse_mean'])} del modelo aprendido "
            f"elegido. El RMSE de la referencia fue {cv_relation} que el del "
            "candidato elegido en validación. El protocolo elige entre siete modelos "
            "aprendidos para contrastarlos con referencias; ese título de "
            "elegido no es una recomendación de despliegue ni afirma superar "
            "todas las referencias en todos los periodos.",
            "## 7. Prueba reservada e interpretación",
            _table(
                ["Predictor", "MAE", "MSE", "RMSE", "R²"],
                [
                    [
                        str(winner["label"])
                        if key == "winner"
                        else "Persistencia de un día"
                        if key == "persistence_1"
                        else "Persistencia semanal",
                        *[
                            _number(_object(_object(value)["metrics"])[metric])
                            for metric in ("mae", "mse", "rmse", "r2")
                        ],
                    ]
                    for key, value in test.items()
                ],
            ),
        ]
    )
    winner_rmse = cast(float, metrics["rmse"])
    points = _objects(winner_test["predictions"])
    largest_error = max(
        points,
        key=lambda point: abs(cast(float, point["actual"]) - cast(float, point["predicted"])),
    )
    parts.append(
        "El RMSE de prueba es mayor que el de validación, lo que limita la "
        "transferencia entre periodos. Como ejemplo concreto, el "
        f"{largest_error['date']} el registro varió {_number(largest_error['actual'], 0)} "
        f"casos y el elegido predijo {_number(largest_error['predicted'], 2)}. "
        "Esa desviación contribuye al error cuadrático, pero el CSV por sí "
        "solo no permite atribuirla a un brote, retraso o corrección. La "
        "diferencia pequeña de validación entre Ridge y SVR lineal tampoco "
        "demuestra superioridad estadística del elegido."
    )
    for key, name in (
        ("persistence_1", "persistencia de un día"),
        ("persistence_7", "persistencia semanal"),
    ):
        baseline_rmse = cast(float, _object(_object(test[key])["metrics"])["rmse"])
        relation = "menor" if winner_rmse < baseline_rmse else "mayor o igual"
        parts.append(
            f"El RMSE del modelo elegido es {relation} que el de {name}: "
            f"{winner_rmse:.4f} frente a {baseline_rmse:.4f}. Esta comparación "
            "describe el bloque reservado; no demuestra que la relación se "
            "mantenga en otra ola, otro país o una revisión distinta de los datos."
        )
    parts.append(
        f"El elegido produjo {winner_test['negative_prediction_count']} predicciones "
        "negativas en prueba. Se conservaron sin recorte. El modelo estima una "
        "variación de registro y no impone una distribución de conteos. Para "
        "exigir salidas no negativas habría que definir otro experimento y "
        "reconocer que esta prueba ya fue observada."
    )
    parts.extend(
        [
            f"![Figura 2. Validación y predicciones de prueba.]"
            f"({Path(os.path.relpath(images[1], report_directory)).as_posix()})",
            "Figura 2. RMSE de validación y comparación observada/predicha del elegido en prueba.",
            "## 8. Parámetros elegidos y avisos de entrenamiento",
            _table(
                ["Modelo", "Parámetros seleccionados"],
                [
                    [str(model["label"]), json.dumps(model["best_params"], ensure_ascii=False)]
                    for model in models
                ],
            ),
            "Una grilla vacía corresponde a un candidato fijo; no significa "
            "que el estimador carezca de parámetros. El JSON conserva además "
            "estimator_parameters, los parámetros completos del estimador "
            "reajustado. fit_warnings contiene categoría y mensaje de los "
            "avisos capturados durante búsqueda y reajuste; su ausencia no "
            "constituye una demostración independiente de optimalidad.",
        ]
    )
    for model in models:
        warnings = _objects(model["fit_warnings"])
        if warnings:
            categories = ", ".join(sorted({str(item["category"]) for item in warnings}))
            parts.append(
                f"{model['label']}: {len(warnings)} avisos ({categories}). "
                "El registro informa que L-BFGS alcanzó el límite de 2000 "
                f"iteraciones en {len(warnings)} ajustes. No se confirmó convergencia "
                "en esos ajustes. Se conservaron configuración y mensajes "
                "originales en report.json; no se aumentó el límite después "
                "de observar la prueba. Los avisos capturados no identifican "
                "por separado el fold de origen."
            )
    parts.extend(
        [
            "## 9. Arquitectura y reproducción",
            "PolarsCovidRepository valida y agrega la fuente. "
            "CovidPreprocessingService construye las entradas pasadas. "
            "build_temporal_split crea las ventanas cronológicas. "
            "covid_model_catalog reutiliza los seis modelos y añade la RNA. "
            "CovidExperimentService usa tune_model, compartido con vinos, "
            "para búsqueda y reajuste, y controla qué modelos acceden a prueba. "
            "La consola exporta series.csv, supervised.csv y report.json. "
            "El visor y este informe consumen evidencia persistida: no entrenan.",
            "Desde la raíz del repositorio, con Python 3.12 o superior y uv:",
            "```powershell\n"
            "uv sync --locked --extra dev\n"
            "uv run covid-quality preprocess\n"
            "uv run covid-quality run\n"
            "uv run python scripts/build_dashboard.py --report outputs/04_covid/report.json "
            "--output outputs/04_covid/results.html\n"
            "```",
            "El resultado histórico entregado ya existe: no es necesario repetir "
            "el entrenamiento para abrir el visor. La consola rechaza carpetas "
            "que ya contienen report.json. Para otra ejecución utiliza "
            "--output outputs/covid_nueva_ejecucion y declara la reutilización "
            "de prueba. --country permite otro nombre exacto del CSV, con "
            "resultados e interpretación propios.",
            "Para reconstruir el Markdown, las figuras y el Word desde los "
            "resultados, instala el extra reports y ejecuta el generador:",
            "```powershell\n"
            "uv sync --locked --extra dev --extra reports\n"
            "uv run --extra reports python scripts/build_covid_report.py\n"
            "```",
            "En equipos sin AVX2, añade --extra compat a sync y a run para "
            "usar el runtime compatible de Polars. No cambia las operaciones "
            "ni el protocolo; sigue fijada la misma versión de la biblioteca.",
            "## 10. Evidencia de ejecución y límites",
            "Las pruebas comprueban agregación y deduplicación geográfica, "
            "retardos, medias pasadas, revisiones negativas, fechas inválidas, "
            "cronología, escalado de X e y por fold, selección por validación "
            "y reconstrucción de métricas desde predicciones. Las pruebas "
            "de vinos permanecen en la suite. Los conteos y resultados de "
            "la revisión final se registran en CHANGELOG.md.",
            "Versiones registradas: "
            + "; ".join(
                f"{name} {value}" for name, value in _object(experiment["versions"]).items()
            )
            + f". La evaluación registró {_number(report['total_seconds'], 2)} segundos "
            "en este equipo; es tiempo de evaluación y no incluye instalación, "
            "carga/exportación ni autoría del informe.",
            f"Los buffers de entradas ocupan {_object(data['matrix_bytes'])['features']} "
            f"bytes y los de respuesta {_object(data['matrix_bytes'])['target']} bytes. "
            "Esto excluye tablas, objetos, copias, árboles y cachés; no representa "
            "el pico de RAM del proceso ni demuestra ahorro por usar Polars.",
            "La copia termina en mayo de 2021 y no describe la situación actual. "
            "La cronología evita usar respuestas futuras al ajustar, pero no "
            "resuelve el problema de revisiones retrospectivas: faltan versiones "
            "del archivo tal como se publicaron cada día. Tampoco se incorporan "
            "población, cantidad de pruebas, vacunación, intervenciones o "
            "mortalidad. No se atribuyen causas ni se emiten recomendaciones sanitarias.",
            "Este experimento sigue siendo regresión. Una matriz de confusión "
            "requiere una tarea y clases justificadas en otro protocolo. La "
            "primera prueba del caso 04 ya es conocida tras esta ejecución; "
            "cualquier ajuste posterior debe reconocerlo y buscar evidencia "
            "nueva para una evaluación confirmatoria.",
            "## 11. Fuentes localizadas",
            "[1] JHU CSSE. README de series temporales, apartado Time series summary. "
            "https://github.com/CSSEGISandData/COVID-19/blob/master/"
            "csse_covid_19_data/csse_covid_19_time_series/README.md",
            "[2] JHU CSSE. README general, apartado Terms of Use. "
            "https://github.com/CSSEGISandData/COVID-19#terms-of-use",
            "[3] JHU CSSE. README de datos, apartado Data modification records. "
            "https://github.com/CSSEGISandData/COVID-19/blob/master/"
            "csse_covid_19_data/README.md#data-modification-records",
            "[4] scikit-learn. TimeSeriesSplit, descripción y Parameters. "
            "https://scikit-learn.org/stable/modules/generated/"
            "sklearn.model_selection.TimeSeriesSplit.html",
            "[5] scikit-learn. TransformedTargetRegressor, descripción y Parameters. "
            "https://scikit-learn.org/stable/modules/generated/"
            "sklearn.compose.TransformedTargetRegressor.html",
            "Consulta de fuentes: 30 de septiembre de 2026. Las elecciones "
            "de país, horizonte, entradas, grillas y prueba pertenecen al "
            "protocolo propio docs/decisions/004-covid-temporal.md.",
        ]
    )
    return "\n\n".join(parts) + "\n"


def _word(markdown: str, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.left_margin, section.right_margin = Cm(2.2), Cm(2.2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor.from_string("23688E")
    header = section.header.paragraphs[0]
    header.text = "Grupo 15 · COVID · Evaluación temporal"
    header.style = document.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.add_run("Caso 04 · Página ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for block in markdown.strip().split("\n\n"):
        if block.startswith("# "):
            document.add_heading(block[2:], level=0)
        elif block.startswith("## "):
            document.add_heading(block[3:], level=1)
        elif block.startswith("!["):
            match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", block)
            if match is None:
                raise ValueError("Imagen Markdown inválida.")
            document.add_picture(str((path.parent / match[2]).resolve()), width=Cm(16.3))
        elif block.startswith("| "):
            rows = [
                [cell.strip() for cell in line.strip("|").split("|")] for line in block.splitlines()
            ]
            table = document.add_table(rows=1, cols=len(rows[0]))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, rows[0], strict=True):
                cell.text = text
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for row in rows[2:]:
                cells = table.add_row().cells
                for cell, text in zip(cells, row, strict=True):
                    cell.text = text
                prevent_split = OxmlElement("w:cantSplit")
                cells[0]._tc.getparent().get_or_add_trPr().append(prevent_split)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
        elif block.startswith("```"):
            commands = block.splitlines()[1:-1]
            for command in commands:
                paragraph = document.add_paragraph(command)
                paragraph.paragraph_format.space_after = Pt(3)
                for run in paragraph.runs:
                    run.font.name, run.font.size = "Consolas", Pt(8)
        else:
            document.add_paragraph(block)
    document.core_properties.title = "Informe integrado de COVID"
    document.core_properties.subject = "Pronóstico temporal de variaciones de casos registrados"
    document.core_properties.author = "Grupo 15"
    document.save(str(path))


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Construir informe COVID desde resultados guardados."
    )
    parser.add_argument("--report", type=Path, default=Path("outputs/04_covid/report.json"))
    parser.add_argument("--output", type=Path, default=Path("docs/reports"))
    arguments = parser.parse_args()
    report = cast(JsonObject, json.loads(arguments.report.read_text(encoding="utf-8")))
    if report.get("schema_version") != "covid-1.0":
        raise ValueError("Este generador requiere un informe covid-1.0.")
    arguments.output.mkdir(parents=True, exist_ok=True)
    images = _figures(report, arguments.report.parent)
    markdown = _markdown(report, images, arguments.output.resolve())
    (arguments.output / "informe_integrado_covid.md").write_text(markdown, encoding="utf-8")
    _word(markdown, arguments.output / "Informe integrado de COVID.docx")
    print(f"Informe Markdown, Word y figuras: {arguments.output.resolve()}")


if __name__ == "__main__":
    main()
